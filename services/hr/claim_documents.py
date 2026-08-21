"""Expense claim document builders — Excel (openpyxl) and PDF (reportlab),
both using the system tenant's letterhead (name/address/contact/UEN).
"""
import io
import os

from django.conf import settings

BLUE = '1F3864'
ORANGE = 'C45911'

CATEGORY_LABELS = {
    'transport': 'Transport', 'meals': 'Meals', 'material': 'Material',
    'equipment': 'Equipment', 'accommodation': 'Accommodation', 'misc': 'Miscellaneous',
}

ITEM_HEADERS = ['Date', 'Category', 'Description', 'Remark', 'Project', 'Amount']


def _tenant_lines(tenant):
    name = (tenant.name if tenant else None) or 'Astronic Svc/Trad Pte Ltd'
    lines = []
    if tenant:
        if tenant.address:
            lines += [ln.strip() for ln in tenant.address.splitlines() if ln.strip()]
        if tenant.phone:
            lines.append(f"Tel: {tenant.phone}")
        if tenant.email:
            lines.append(f"Email: {tenant.email}")
        if tenant.uen:
            lines.append(f"UEN: {tenant.uen}")
    return name, lines


def _logo_path(tenant):
    if tenant and tenant.logo:
        path = os.path.join(settings.MEDIA_ROOT, str(tenant.logo))
        if os.path.exists(path):
            return path
    return None


def _item_row(it):
    return [
        it.expense_date.strftime('%d/%m/%Y') if it.expense_date else '',
        CATEGORY_LABELS.get(it.category, it.category),
        it.description or '',
        it.remark or '',
        it.project_no or '',
        float(it.amount or 0),
    ]


# ── Excel ─────────────────────────────────────────────────────────────────
def build_excel(claim, tenant=None):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    name, addr_lines = _tenant_lines(claim.claimant.tenant if hasattr(claim.claimant, 'tenant') else tenant)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Claim'

    bold = Font(bold=True)
    header_font = Font(bold=True, size=14, color=BLUE)
    grey = Font(color='595959', size=9)

    row = 1
    ws.cell(row=row, column=1, value=name).font = header_font
    row += 1
    for ln in addr_lines:
        ws.cell(row=row, column=1, value=ln).font = grey
        row += 1
    row += 1

    ws.cell(row=row, column=1, value='EXPENSE CLAIM').font = Font(bold=True, size=12)
    row += 2

    meta = [
        ('Title', claim.title),
        ('Period', claim.period_month.strftime('%B %Y') if claim.period_month else ''),
        ('Claimant', claim.claimant.full_name if hasattr(claim.claimant, 'full_name') else str(claim.claimant)),
        ('Status', claim.get_status_display()),
        ('Submitted', claim.submitted_at.strftime('%d %B %Y, %H:%M') if claim.submitted_at else ''),
        ('Approver', getattr(claim.approver, 'full_name', None) or (str(claim.approver) if claim.approver else '')),
    ]
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    header_row = row
    for col, h in enumerate(ITEM_HEADERS, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
    row += 1

    for it in claim.items.all():
        for col, val in enumerate(_item_row(it), start=1):
            c = ws.cell(row=row, column=col, value=val)
            if col == 6:
                c.number_format = '#,##0.00'
        row += 1

    ws.cell(row=row, column=5, value='TOTAL').font = bold
    total_cell = ws.cell(row=row, column=6, value=float(claim.total_amount or 0))
    total_cell.font = bold
    total_cell.number_format = '#,##0.00'

    if claim.notes:
        row += 2
        ws.cell(row=row, column=1, value='Notes').font = bold
        row += 1
        ws.cell(row=row, column=1, value=claim.notes)

    if claim.remarks:
        row += 2
        ws.cell(row=row, column=1, value='Approver remarks').font = bold
        row += 1
        ws.cell(row=row, column=1, value=claim.remarks)

    widths = [12, 14, 36, 28, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── PDF ───────────────────────────────────────────────────────────────────
def build_pdf(claim, tenant=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage

    BLUE_C = colors.HexColor(f'#{BLUE}')
    ORANGE_C = colors.HexColor(f'#{ORANGE}')

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('Addr', parent=styles['Normal'], fontName='Helvetica', fontSize=8, textColor=colors.HexColor('#555555')))
    styles.add(ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=10))
    styles.add(ParagraphStyle('Label', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, spaceBefore=4, spaceAfter=1))

    tenant = claim.claimant.tenant if hasattr(claim.claimant, 'tenant') else tenant
    name, addr_lines = _tenant_lines(tenant)

    story = []

    logo = _logo_path(tenant)
    left = []
    if logo:
        try:
            from PIL import Image as PILImage
            iw, ih = PILImage.open(logo).size
            w = 32 * mm
            left.append(RLImage(logo, width=w, height=w * ih / iw))
        except Exception:
            pass
    left.append(Paragraph('EXPENSE CLAIM', ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=14, textColor=BLUE_C)))
    right = [Paragraph(f"<b>{name}</b>", styles['Body'])] + [Paragraph(ln, styles['Addr']) for ln in addr_lines]
    header = Table([[left, right]], colWidths=[95 * mm, 80 * mm])
    header.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                ('LINEBELOW', (0, 0), (-1, -1), 1.5, ORANGE_C)]))
    story += [header, Spacer(1, 8)]

    claimant_name = getattr(claim.claimant, 'full_name', None) or str(claim.claimant)
    approver_name = getattr(claim.approver, 'full_name', None) or (str(claim.approver) if claim.approver else '—')
    meta_left = [
        Paragraph(f"<b>Title:</b> {claim.title}", styles['Body']),
        Paragraph(f"<b>Period:</b> {claim.period_month.strftime('%B %Y') if claim.period_month else ''}", styles['Body']),
        Paragraph(f"<b>Claimant:</b> {claimant_name}", styles['Body']),
    ]
    meta_right = [
        Paragraph(f"<b>Status:</b> {claim.get_status_display()}", styles['Body']),
        Paragraph(f"<b>Submitted:</b> {claim.submitted_at.strftime('%d %b %Y, %H:%M') if claim.submitted_at else '—'}", styles['Body']),
        Paragraph(f"<b>Approver:</b> {approver_name}", styles['Body']),
    ]
    meta = Table([[meta_left, meta_right]], colWidths=[95 * mm, 80 * mm])
    meta.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    story += [meta, Spacer(1, 10)]

    data = [ITEM_HEADERS]
    for it in claim.items.all():
        row = _item_row(it)
        row[5] = f"{row[5]:,.2f}"
        data.append([Paragraph(str(v), styles['Body']) if i in (2, 3) else v for i, v in enumerate(row)])
    data.append(['', '', '', '', 'TOTAL', f"{float(claim.total_amount or 0):,.2f}"])

    tbl = Table(data, colWidths=[20 * mm, 22 * mm, 48 * mm, 40 * mm, 22 * mm, 23 * mm])
    tbl.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, 0), (-1, 0), BLUE_C),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (5, 0), (5, -1), 'RIGHT'),
        ('LINEBELOW', (0, 0), (-1, 0), 1, ORANGE_C),
        ('SPAN', (0, -1), (4, -1)),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, ORANGE_C),
        ('ALIGN', (4, -1), (4, -1), 'RIGHT'),
    ]))
    story.append(tbl)

    if claim.notes:
        story += [Spacer(1, 8), Paragraph('Notes:', styles['Label']), Paragraph(claim.notes, styles['Body'])]
    if claim.remarks:
        story += [Spacer(1, 8), Paragraph('Approver remarks:', styles['Label']), Paragraph(claim.remarks, styles['Body'])]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
                             leftMargin=14 * mm, rightMargin=14 * mm)
    doc.build(story)
    buf.seek(0)
    return buf
